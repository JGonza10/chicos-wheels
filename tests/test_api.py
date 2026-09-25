"""
Pruebas de la API contra una base de datos temporal.

Verifican las reglas de negocio que más duelen si fallan:
el dinero, el stock y el aislamiento entre cuentas.

    python -m unittest discover -s tests -v
"""
import io
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image

# La base de prueba debe definirse ANTES de importar la app.
_TMP = Path(tempfile.gettempdir()) / f"ch-test-{os.getpid()}.db"
os.environ["DB_FILE"] = str(_TMP)
os.environ["JWT_SECRET"] = "llave-de-prueba"
os.environ["COLLECTHUB_SIN_RESPALDO"] = "1"

from collecthub import crear_app  # noqa: E402


class PruebasAPI(unittest.TestCase):
    """Se ejecutan en orden alfabético: los nombres llevan número a propósito."""

    @classmethod
    def setUpClass(cls):
        cls.app = crear_app()
        cls.c = cls.app.test_client()
        cls.token = None
        cls.token2 = None
        cls.articulo = None
        cls.plataforma = None

    @classmethod
    def tearDownClass(cls):
        for sufijo in ("", "-wal", "-shm"):
            Path(str(_TMP) + sufijo).unlink(missing_ok=True)
        import shutil
        shutil.rmtree(_TMP.parent / "fotos", ignore_errors=True)

    # --- ayudantes ---
    def pedir(self, metodo, ruta, cuerpo=None, token="propio"):
        tk = self.token if token == "propio" else token
        cab = {"Authorization": f"Bearer {tk}"} if tk else {}
        r = getattr(self.c, metodo.lower())(ruta, json=cuerpo, headers=cab)
        return r.status_code, (r.get_json() or {})

    # ---------------- Cuentas ----------------

    def test_01_rechaza_datos_invalidos(self):
        s, _ = self.pedir("POST", "/api/auth/registro", {"email": "a@b.co", "password": "123"}, None)
        self.assertEqual(s, 400, "una contraseña de 3 caracteres no debe pasar")
        s, _ = self.pedir("POST", "/api/auth/registro",
                          {"email": "sinarroba", "password": "contrasena1"}, None)
        self.assertEqual(s, 400, "un correo mal escrito no debe pasar")

    def test_02_registro_deja_la_cuenta_lista(self):
        s, d = self.pedir("POST", "/api/auth/registro",
                          {"email": "chicos@wheels.mx", "password": "coleccion2026",
                           "nombre": "Chicos"}, None)
        self.assertEqual(s, 201)
        PruebasAPI.token = d["token"]
        _, plats = self.pedir("GET", "/api/plataformas")
        self.assertEqual(len(plats), 5, "debe traer Mercado Libre, eBay, Facebook, Bazar y Balderas")
        PruebasAPI.plataforma = next(p["id"] for p in plats if p["codigo"] == "ML")

    def test_03_no_repite_correo(self):
        s, _ = self.pedir("POST", "/api/auth/registro",
                          {"email": "chicos@wheels.mx", "password": "otracosa123"}, None)
        self.assertEqual(s, 409)

    def test_04_sin_token_no_hay_acceso(self):
        s, _ = self.pedir("GET", "/api/estado", token=None)
        self.assertEqual(s, 401)

    def test_05_login_no_revela_si_el_correo_existe(self):
        s, d = self.pedir("POST", "/api/auth/login",
                          {"email": "chicos@wheels.mx", "password": "mala"}, None)
        self.assertEqual(s, 401)
        self.assertIn("Correo o contraseña", d["error"])

    # ---------------- Inventario ----------------

    def test_06_crea_pieza_con_valuacion_inicial(self):
        s, d = self.pedir("POST", "/api/articulos", {
            "tipo": "Hot Wheels", "nombre": "Datsun 240Z", "precio_compra": 100,
            "valor_estimado": 500, "cantidad": 4, "fecha_adq": "2026-01-10", "ubicacion": "Caja A",
        })
        self.assertEqual(s, 201)
        PruebasAPI.articulo = d["id"]
        self.assertEqual(d["disponible"], 4)
        _, det = self.pedir("GET", f"/api/articulos/{d['id']}")
        self.assertEqual(len(det["valuaciones"]), 1,
                         "el valor inicial debe quedar en el historial de precios")

    def test_07_rechaza_pieza_sin_nombre(self):
        s, _ = self.pedir("POST", "/api/articulos", {"tipo": "Hot Wheels", "nombre": "   "})
        self.assertEqual(s, 400)

    # ---------------- Dinero ----------------

    def test_08_la_ganancia_la_calcula_la_base_de_datos(self):
        # 500 cobrados − 100 costo − 25 cuota fija − 65 comisión (13%) − 80 envío − 20 otros = 210
        s, d = self.pedir("POST", "/api/ventas", {
            "articulo_id": self.articulo, "plataforma_id": self.plataforma, "cantidad": 1,
            "precio": 500, "envio": 80, "otros": 20, "fecha": "2026-07-20",
            "ganancia_neta": 999999,   # intento de manipulación desde el cliente
        })
        self.assertEqual(s, 201)
        self.assertAlmostEqual(d["ganancia_neta"], 210, places=2,
                               msg="el número que mandó el cliente debe ignorarse")

    def test_09_vender_descuenta_stock(self):
        _, d = self.pedir("GET", f"/api/articulos/{self.articulo}")
        self.assertEqual(d["cantidad"], 3)

    def test_10_apartado_reserva_y_bloquea(self):
        s, ap = self.pedir("POST", "/api/apartados", {
            "articulo_id": self.articulo, "cantidad": 3, "precio_acordado": 600,
            "anticipo": 200, "fecha_limite": "2099-01-01", "cliente_nuevo": "Cliente Diez",
        })
        self.assertEqual(s, 201)

        _, art = self.pedir("GET", f"/api/articulos/{self.articulo}")
        self.assertEqual(art["cantidad"], 3)
        self.assertEqual(art["disponible"], 0, "todo el stock quedó comprometido")

        s, err = self.pedir("POST", "/api/ventas", {
            "articulo_id": self.articulo, "plataforma_id": self.plataforma, "precio": 300})
        self.assertEqual(s, 409, "no debe permitir vender lo que está apartado")
        self.assertIn("apartado", err["error"].lower())

        # Liquidar sí procede y consume el stock reservado.
        s, venta = self.pedir("POST", "/api/ventas", {
            "articulo_id": self.articulo, "plataforma_id": self.plataforma, "cantidad": 3,
            "precio": 600, "apartado_id": ap["id"]})
        self.assertEqual(s, 201)
        self.assertEqual(venta["anticipo_aplicado"], 200)

        _, art = self.pedir("GET", f"/api/articulos/{self.articulo}")
        self.assertEqual(art["cantidad"], 0)
        _, aps = self.pedir("GET", "/api/apartados")
        self.assertEqual(aps[0]["estatus"], "Liquidado")

    def test_11_anticipo_no_supera_el_precio(self):
        _, a = self.pedir("POST", "/api/articulos", {"tipo": "Pokémon", "nombre": "Carta prueba"})
        s, _ = self.pedir("POST", "/api/apartados",
                          {"articulo_id": a["id"], "precio_acordado": 100, "anticipo": 500})
        self.assertEqual(s, 400)

    def test_12_lote_reparte_precio_y_cobra_cuota_una_vez(self):
        _, a1 = self.pedir("POST", "/api/articulos", {
            "tipo": "Hot Wheels", "nombre": "Lote A", "precio_compra": 20, "valor_estimado": 300})
        _, a2 = self.pedir("POST", "/api/articulos", {
            "tipo": "Hot Wheels", "nombre": "Lote B", "precio_compra": 20, "valor_estimado": 100})
        s, d = self.pedir("POST", "/api/ventas/lote", {
            "articulo_ids": [a1["id"], a2["id"]], "plataforma_id": self.plataforma,
            "precio": 400, "envio": 50, "fecha": "2026-07-25"})
        self.assertEqual(s, 201)
        v1, v2 = d["ventas"]
        self.assertAlmostEqual(v1["precio"], 300, places=2, msg="reparte según el valor de mercado")
        self.assertAlmostEqual(v2["precio"], 100, places=2)
        self.assertEqual(v1["com_fija"] + v2["com_fija"], 25, "la cuota fija se cobra una vez")
        self.assertEqual(v1["envio"] + v2["envio"], 50, "el envío es del paquete completo")

    def test_13_lote_necesita_dos_piezas(self):
        s, _ = self.pedir("POST", "/api/ventas/lote",
                          {"articulo_ids": ["x"], "plataforma_id": self.plataforma, "precio": 100})
        self.assertEqual(s, 400)

    def test_14_intercambio_mueve_inventario(self):
        _, a = self.pedir("POST", "/api/articulos", {
            "tipo": "Hot Wheels", "nombre": "Bel Air 57", "valor_estimado": 180})
        s, d = self.pedir("POST", "/api/intercambios", {
            "contraparte": "Puesto 14", "entrega_ids": [a["id"]],
            "recibidos": [{"nombre": "Skyline C210", "tipo": "Hot Wheels", "valor": 420}]})
        self.assertEqual(s, 201)
        self.assertEqual(d["diferencia"], 240, "420 recibidos menos 180 entregados")

        _, est = self.pedir("GET", "/api/estado")
        entregada = next(x for x in est["articulos"] if x["id"] == a["id"])
        self.assertEqual(entregada["cantidad"], 0, "la entregada sale del inventario")
        recibida = next((x for x in est["articulos"] if x["nombre"] == "Skyline C210"), None)
        self.assertIsNotNone(recibida, "la recibida entra al inventario")
        self.assertEqual(recibida["precio_compra"], 0, "entra con costo cero")

    def test_15_cancelar_venta_devuelve_stock(self):
        _, a = self.pedir("POST", "/api/articulos", {
            "tipo": "Hot Wheels", "nombre": "Devolución", "cantidad": 2,
            "precio_compra": 10, "valor_estimado": 50})
        _, v = self.pedir("POST", "/api/ventas", {
            "articulo_id": a["id"], "plataforma_id": self.plataforma, "cantidad": 2, "precio": 100})
        _, d = self.pedir("GET", f"/api/articulos/{a['id']}")
        self.assertEqual(d["cantidad"], 0)
        self.pedir("DELETE", f"/api/ventas/{v['id']}")
        _, d = self.pedir("GET", f"/api/articulos/{a['id']}")
        self.assertEqual(d["cantidad"], 2)

    def test_16_cambiar_comision_no_toca_el_historial(self):
        _, antes = self.pedir("GET", "/api/ventas")
        previos = sorted(v["ganancia_neta"] for v in antes)
        self.pedir("PATCH", f"/api/plataformas/{self.plataforma}",
                   {"nombre": "Mercado Libre", "codigo": "ML", "com_pct": 0.30, "com_fija": 99})
        _, despues = self.pedir("GET", "/api/ventas")
        self.assertEqual(previos, sorted(v["ganancia_neta"] for v in despues),
                         "las ventas ya registradas deben quedar congeladas")

    # ---------------- Valuación y métricas ----------------

    def test_17_valuacion_mueve_precio_y_guarda_tendencia(self):
        s, d = self.pedir("POST", f"/api/articulos/{self.articulo}/valuaciones",
                          {"valor": 750, "fuente": "eBay vendidos"})
        self.assertEqual(s, 201)
        self.assertEqual(d["articulo"]["valor_estimado"], 750)
        self.assertGreaterEqual(len(d["valuaciones"]), 2)

    def test_18_metricas_en_sql_coinciden_con_las_ventas(self):
        _, st = self.pedir("GET", "/api/stats")
        _, ventas = self.pedir("GET", "/api/ventas")
        self.assertAlmostEqual(st["ganancia"], sum(v["ganancia_neta"] for v in ventas), places=2)
        self.assertIsInstance(st["por_canal"], list)

    def test_19_precio_objetivo(self):
        # Para ganar 200 en un canal al 13% + 25 fijos, con 45 de costo:
        # (200 + 45 + 25) / 0.87 = 310.34
        s, d = self.pedir("POST", "/api/calculadora/precio-objetivo", {
            "plataforma_id": self.plataforma, "deseado": 200, "costo_total": 45})
        self.assertEqual(s, 200)
        self.assertGreater(d["precio"], 245, "debe cubrir costo, comisión y ganancia")

    # ---------------- Aislamiento ----------------

    def test_20_una_cuenta_no_ve_la_de_otra(self):
        _, d = self.pedir("POST", "/api/auth/registro",
                          {"email": "ajena@wheels.mx", "password": "otracuenta123"}, None)
        PruebasAPI.token2 = d["token"]
        s, est = self.pedir("GET", "/api/estado", token=self.token2)
        self.assertEqual(len(est["articulos"]), 0, "la cuenta nueva arranca vacía")
        s, _ = self.pedir("GET", f"/api/articulos/{self.articulo}", token=self.token2)
        self.assertEqual(s, 404, "no debe poder leer piezas ajenas")
        s, _ = self.pedir("DELETE", f"/api/articulos/{self.articulo}", token=self.token2)
        self.assertEqual(s, 404, "no debe poder borrar piezas ajenas")

    def test_21_datos_de_ejemplo_completos(self):
        s, d = self.pedir("POST", "/api/seed", {}, token=self.token2)
        self.assertEqual(s, 200)
        self.assertEqual(len(d["articulos"]), 7)
        self.assertEqual(len(d["ventas"]), 4)
        self.assertEqual(len(d["apartados"]), 2)
        self.assertEqual(len(d["intercambios"]), 1)
        self.assertTrue(any(len(a["valuaciones"]) >= 3 for a in d["articulos"]),
                        "debe haber historial de precios para ver la tendencia")

    def test_22_apartado_vencido_se_marca_solo(self):
        _, est = self.pedir("GET", "/api/estado", token=self.token2)
        self.assertTrue(any(a["estatus"] == "Vencido" for a in est["apartados"]))

    # ---------------- Foto → IA ----------------
    # identificar_foto()/buscar_precio_pokemon() se mockean: las pruebas no
    # deben depender de red ni gastar la API real de Claude.

    @staticmethod
    def _imagen():
        buf = io.BytesIO()
        Image.new("RGB", (10, 10), (255, 0, 0)).save(buf, format="JPEG")
        buf.seek(0)
        return buf

    def test_23_identificar_exige_sesion(self):
        r = self.c.post("/api/articulos/identificar",
                        data={"foto": (self._imagen(), "foto.jpg")}, content_type="multipart/form-data")
        self.assertEqual(r.status_code, 401)

    def test_24_identificar_rechaza_archivo_que_no_es_imagen(self):
        r = self.c.post("/api/articulos/identificar",
                        data={"foto": (io.BytesIO(b"no es una imagen"), "foto.jpg")},
                        headers={"Authorization": f"Bearer {self.token}"}, content_type="multipart/form-data")
        self.assertEqual(r.status_code, 400)

    def test_25_identificar_regresa_la_sugerencia(self):
        falsa = {"tipo": "Hot Wheels", "confianza": 0.82, "nombre": "Datsun 240Z", "numero": "",
                "serie": "", "color": "Rojo", "expansion": "", "rareza": "", "grado": "", "cert": "",
                "notas": "Auto rojo con base metálica"}
        with patch("collecthub.rutas.articulos.identificar_foto", return_value=falsa):
            r = self.c.post("/api/articulos/identificar",
                            data={"foto": (self._imagen(), "foto.jpg")},
                            headers={"Authorization": f"Bearer {self.token}"}, content_type="multipart/form-data")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()["nombre"], "Datsun 240Z")

    def test_26_identificar_pokemon_agrega_precio_de_mercado(self):
        falsa = {"tipo": "Pokémon", "confianza": 0.91, "nombre": "Charizard ex", "numero": "",
                "serie": "", "color": "", "expansion": "Obsidian Flames", "rareza": "",
                "grado": "", "cert": "", "notas": ""}
        precio = {"expansion": "Obsidian Flames", "numero": "125", "rareza": "Ultra Rare", "valor_estimado": 45.5}
        with patch("collecthub.rutas.articulos.identificar_foto", return_value=falsa), \
             patch("collecthub.rutas.articulos.buscar_precio_pokemon", return_value=precio):
            r = self.c.post("/api/articulos/identificar",
                            data={"foto": (self._imagen(), "foto.jpg")},
                            headers={"Authorization": f"Bearer {self.token}"}, content_type="multipart/form-data")
        d = r.get_json()
        self.assertEqual(d["rareza"], "Ultra Rare")
        self.assertEqual(d["valor_estimado"], 45.5)

    def test_27_foto_guardada_queda_aislada_entre_cuentas(self):
        _, art = self.pedir("POST", "/api/articulos", {"tipo": "Hot Wheels", "nombre": "Prueba de foto"})
        r = self.c.post(f"/api/articulos/{art['id']}/foto",
                        data={"foto": (self._imagen(), "foto.jpg")},
                        headers={"Authorization": f"Bearer {self.token}"}, content_type="multipart/form-data")
        self.assertEqual(r.status_code, 200)
        archivo = r.get_json()["foto"][len("local:"):]

        r = self.c.get(f"/api/articulos/foto/{archivo}", headers={"Authorization": f"Bearer {self.token}"})
        self.assertEqual(r.status_code, 200, "el dueño sí debe poder ver su propia foto")

        r = self.c.get(f"/api/articulos/foto/{archivo}", headers={"Authorization": f"Bearer {self.token2}"})
        self.assertEqual(r.status_code, 404, "otra cuenta no debe poder ver esta foto ni adivinando el nombre")

        r = self.c.get(f"/api/articulos/foto/{archivo}")
        self.assertEqual(r.status_code, 401, "sin sesión, tampoco")

    def test_28_foto_rechaza_nombre_de_archivo_invalido(self):
        r = self.c.get("/api/articulos/foto/..%2f..%2fetc%2fpasswd",
                       headers={"Authorization": f"Bearer {self.token}"})
        self.assertIn(r.status_code, (400, 404))

    def test_29_reporte_pdf_configurable(self):
        h = {"Authorization": f"Bearer {self.token}"}
        r = self.c.post("/api/reporte-pdf", json={}, headers=h)
        self.assertEqual(r.status_code, 400, "sin secciones no hay reporte")
        cfg = {"secciones": ["resumen", "inventario", "ventas", "apartados"], "estatus": "todos",
               "orientacion": "horizontal", "titulo": "Prueba – Pokémon"}
        r = self.c.post("/api/reporte-pdf", json=cfg, headers=h)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data.startswith(b"%PDF"))
        for est in ("disponible", "apartado", "conservar", "agotado", "estancado"):
            r = self.c.post("/api/reporte-pdf", json={**cfg, "estatus": est, "tipo": "Pokémon"}, headers=h)
            self.assertEqual(r.status_code, 200, est)
        r = self.c.post("/api/reporte-pdf", json={**cfg, "estatus": "inventado"}, headers=h)
        self.assertEqual(r.status_code, 400)
        self.assertEqual(self.c.post("/api/reporte-pdf", json=cfg).status_code, 401)

    def test_30_fecha_de_entrega_se_guarda_y_se_limpia(self):
        h = {"Authorization": f"Bearer {self.token}"}
        _, art = self.pedir("POST", "/api/articulos", {"tipo": "Hot Wheels", "nombre": "Entrega", "cantidad": 2,
                                                       "precio_compra": 10, "valor_estimado": 30})
        _, plat = self.pedir("GET", "/api/estado")
        pid = plat["plataformas"][0]["id"]
        _, v = self.pedir("POST", "/api/ventas", {"articulo_id": art["id"], "plataforma_id": pid,
                                                  "cantidad": 1, "precio": 30, "envio": 5})
        vid = v["id"]
        r = self.c.patch(f"/api/ventas/{vid}", json={"estatus_envio": "Entregado"}, headers=h)
        self.assertEqual(r.status_code, 200)
        self.assertRegex(r.get_json()["fecha_entrega"], r"^\d{4}-\d{2}-\d{2}$")
        r = self.c.patch(f"/api/ventas/{vid}", json={"estatus_envio": "Entregado", "fecha_entrega": "2026-01-15"}, headers=h)
        self.assertEqual(r.get_json()["fecha_entrega"], "2026-01-15")
        r = self.c.patch(f"/api/ventas/{vid}", json={"estatus_envio": "Enviado"}, headers=h)
        self.assertEqual(r.get_json()["fecha_entrega"], "", "al regresar de Entregado la fecha se borra")
        r = self.c.get(f"/api/ventas/{vid}/recibo", headers=h)
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data.startswith(b"%PDF"))
        r = self.c.get(f"/api/ventas/{vid}/recibo", headers={"Authorization": f"Bearer {self.token2}"})
        self.assertEqual(r.status_code, 404, "otra cuenta no puede bajar el recibo")

    def test_31_pedidos_de_clientes(self):
        h = {"Authorization": f"Bearer {self.token}"}
        h2 = {"Authorization": f"Bearer {self.token2}"}
        self.assertEqual(self.c.post("/api/pedidos", json={}, headers=h).status_code, 400)
        r = self.c.post("/api/pedidos", json={"descripcion": "Datsun 240Z verde", "tope": 150}, headers=h)
        self.assertEqual(r.status_code, 201)
        pid = r.get_json()["id"]
        self.assertEqual(self.c.patch(f"/api/pedidos/{pid}", json={"atendido": True}, headers=h).get_json()["atendido"], 1)
        self.assertEqual(self.c.patch(f"/api/pedidos/{pid}", json={"atendido": True}, headers=h2).status_code, 404)
        estado = self.c.get("/api/estado", headers=h).get_json()
        self.assertTrue(any(p["id"] == pid for p in estado["pedidos"]))
        self.assertEqual(self.c.delete(f"/api/pedidos/{pid}", headers=h).status_code, 200)

    def test_32_foto_lista_para_publicar(self):
        h = {"Authorization": f"Bearer {self.token}"}
        _, art = self.pedir("POST", "/api/articulos", {"tipo": "Hot Wheels", "nombre": "Foto FB"})
        r = self.c.get(f"/api/articulos/{art['id']}/foto-publicar", headers=h)
        self.assertEqual(r.status_code, 400, "sin foto propia no hay nada que preparar")
        self.c.post(f"/api/articulos/{art['id']}/foto", data={"foto": (self._imagen(), "f.jpg")},
                    headers=h, content_type="multipart/form-data")
        r = self.c.get(f"/api/articulos/{art['id']}/foto-publicar", headers=h)
        self.assertEqual(r.status_code, 200)
        from PIL import Image
        self.assertEqual(Image.open(io.BytesIO(r.data)).size, (1080, 1080))
        r = self.c.get(f"/api/articulos/{art['id']}/foto-publicar", headers={"Authorization": f"Bearer {self.token2}"})
        self.assertEqual(r.status_code, 404)

    def test_33_datos_desde_link_de_mattel(self):
        h = {"Authorization": f"Bearer {self.token}"}
        ficha = json.dumps({"title": "Hot Wheels Demo", "price": 2500, "vendor": "Hot Wheels Collectors",
                            "available": True, "featured_image": "//cdn.shopify.com/x.jpg", "images": [],
                            "description": "<p>Pieza <b>demo</b></p>"}).encode()
        with patch("collecthub.mattel._leer", return_value=ficha), patch("collecthub.mattel.tipo_de_cambio", return_value=20.0):
            r = self.c.post("/api/articulos/desde-mattel", json={"url": "https://creations.mattel.com/products/demo-hcd19"}, headers=h)
            self.assertEqual(r.status_code, 200)
            d = r.get_json()
            self.assertEqual((d["nombre"], d["precio_usd"], d["precio_mxn"]), ("Hot Wheels Demo", 25.0, 500.0))
            self.assertEqual(d["imagen"], "https://cdn.shopify.com/x.jpg")
            self.assertEqual(d["descripcion"], "Pieza demo")
        for malo in ("http://creations.mattel.com/products/x", "https://evil.com/products/x",
                     "https://mattel.com.evil.com/products/x", "https://creations.mattel.com/collections/x",
                     "https://localhost:3000/products/x", ""):
            r = self.c.post("/api/articulos/desde-mattel", json={"url": malo}, headers=h)
            self.assertEqual(r.status_code, 400, malo)
        self.assertEqual(self.c.post("/api/articulos/desde-mattel", json={"url": "https://creations.mattel.com/products/x"}).status_code, 401)

    def test_34_respaldo_diario_conserva_solo_los_ultimos(self):
        from collecthub import respaldo
        with tempfile.TemporaryDirectory() as d, patch.object(respaldo, "carpeta", return_value=Path(d)):
            for dia in range(1, 21):
                (Path(d) / f"collecthub-2026-01-{dia:02d}.db").write_bytes(b"x")
            hoy_ = respaldo.hacer_respaldo()
            self.assertTrue(hoy_ and hoy_.exists())
            self.assertIsNone(respaldo.hacer_respaldo(), "un solo respaldo por día")
            self.assertEqual(len(list(Path(d).glob("collecthub-*.db"))), respaldo.CONSERVAR)
            import sqlite3
            con = sqlite3.connect(hoy_)
            try:
                self.assertTrue(con.execute("select count(*) from usuarios").fetchone())
            finally:
                con.close()

    def _pieza(self, nombre, cantidad, compra, valor):
        _, a = self.pedir("POST", "/api/articulos", {"tipo": "Hot Wheels", "nombre": nombre, "cantidad": cantidad,
                                                     "precio_compra": compra, "valor_estimado": valor,
                                                     "ubicacion": "Caja T"})
        return a

    def _libres(self, id_art):
        _, est = self.pedir("GET", "/api/estado")
        return next(a["disponible"] for a in est["articulos"] if a["id"] == id_art)

    def test_35_encargo_reserva_stock_y_calcula_costos(self):
        h = {"Authorization": f"Bearer {self.token}"}
        a1, a2 = self._pieza("Enc A", 3, 40, 100), self._pieza("Enc B", 1, 10, 30)
        r = self.c.post("/api/encargos", headers=h, json={
            "cliente_nuevo": "Cliente Balderas", "tel_nuevo": "555", "fecha_entrega": "2026-10-03",
            "anticipo": 100, "forma_anticipo": "Depósito",
            "items": [{"articulo_id": a1["id"], "cantidad": 2}, {"articulo_id": a2["id"], "cantidad": 1, "precio_unit": 35}]})
        self.assertEqual(r.status_code, 201, r.get_json())
        e = r.get_json()
        self.assertEqual((e["piezas"], e["total"], e["costo"], e["ganancia"], e["resta"]), (3, 235.0, 90.0, 145.0, 135.0))
        self.assertEqual(self._libres(a1["id"]), 1, "el encargo reserva stock")
        self.assertEqual(self._libres(a2["id"]), 0)
        # no se puede encargar lo que ya está reservado
        r = self.c.post("/api/encargos", headers=h, json={"comprador_id": e["comprador_id"],
                        "items": [{"articulo_id": a2["id"], "cantidad": 1}]})
        self.assertEqual(r.status_code, 409)
        # ni un anticipo mayor al total, ni sin piezas, ni sin cliente
        self.assertEqual(self.c.post("/api/encargos", headers=h, json={"comprador_id": e["comprador_id"], "anticipo": 9999,
                         "items": [{"articulo_id": a1["id"], "cantidad": 1}]}).status_code, 400)
        self.assertEqual(self.c.post("/api/encargos", headers=h, json={"comprador_id": e["comprador_id"], "items": []}).status_code, 400)
        self.assertEqual(self.c.post("/api/encargos", headers=h, json={"items": [{"articulo_id": a1["id"], "cantidad": 1}]}).status_code, 400)
        # la pieza no se puede borrar mientras esté en un encargo pendiente
        self.assertEqual(self.c.delete(f"/api/articulos/{a1['id']}", headers=h).status_code, 409)
        # editar (PUT) puede volver a pedir lo que el mismo encargo tenía reservado
        r = self.c.put(f"/api/encargos/{e['id']}", headers=h, json={"comprador_id": e["comprador_id"], "anticipo": 0,
                       "items": [{"articulo_id": a1["id"], "cantidad": 3}]})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(self._libres(a1["id"]), 0)
        self.assertEqual(self._libres(a2["id"]), 1, "al quitarla del encargo vuelve a estar libre")
        # aislamiento: otra cuenta no ve ni toca este encargo
        h2 = {"Authorization": f"Bearer {self.token2}"}
        self.assertEqual(self.c.patch(f"/api/encargos/{e['id']}", headers=h2, json={"estatus": "Cancelado"}).status_code, 404)
        # cancelar libera el stock
        r = self.c.patch(f"/api/encargos/{e['id']}", headers=h, json={"estatus": "Cancelado"})
        self.assertEqual(r.get_json()["estatus"], "Cancelado")
        self.assertEqual(self._libres(a1["id"]), 3)
        self.assertEqual(self.c.delete(f"/api/encargos/{e['id']}", headers=h).status_code, 200)

    def test_36_entregar_encargo_lo_convierte_en_ventas(self):
        h = {"Authorization": f"Bearer {self.token}"}
        a1, a2 = self._pieza("Ent A", 2, 50, 120), self._pieza("Ent B", 1, 20, 60)
        _, est = self.pedir("GET", "/api/estado")
        ventas_antes = len(est["ventas"])
        e = self.c.post("/api/encargos", headers=h, json={
            "cliente_nuevo": "Comprador Entrega", "anticipo": 50, "forma_anticipo": "Efectivo",
            "items": [{"articulo_id": a1["id"], "cantidad": 2}, {"articulo_id": a2["id"], "cantidad": 1}]}).get_json()
        self.assertEqual(e["total"], 300.0)
        r = self.c.patch(f"/api/encargos/{e['id']}", headers=h, json={"estatus": "Empacado"})
        self.assertEqual(r.get_json()["estatus"], "Empacado")
        r = self.c.post(f"/api/encargos/{e['id']}/entregar", headers=h, json={"cobrado": 200, "forma": "Depósito"})
        self.assertEqual(r.status_code, 200, r.get_json())
        d = r.get_json()
        self.assertEqual((d["estatus"], d["debe"], d["cobrado_entrega"]), ("Entregado", 50.0, 200.0))
        _, est = self.pedir("GET", "/api/estado")
        nuevas = est["ventas"][:len(est["ventas"]) - ventas_antes]
        nuevas = [v for v in est["ventas"] if v["nombre_snap"] in ("Ent A", "Ent B")]
        self.assertEqual(len(nuevas), 2)
        self.assertAlmostEqual(sum(v["precio"] for v in nuevas), 300.0)
        self.assertAlmostEqual(sum(v["ganancia_neta"] for v in nuevas), 300 - (2 * 50 + 20))
        self.assertEqual(next(a["cantidad"] for a in est["articulos"] if a["id"] == a1["id"]), 0)
        # ya entregado: no se entrega dos veces, no se cancela, no se borra
        self.assertEqual(self.c.post(f"/api/encargos/{e['id']}/entregar", headers=h, json={}).status_code, 409)
        self.assertEqual(self.c.patch(f"/api/encargos/{e['id']}", headers=h, json={"estatus": "Cancelado"}).status_code, 409)
        self.assertEqual(self.c.delete(f"/api/encargos/{e['id']}", headers=h).status_code, 409)

    def test_37_hoja_de_entrega_pdf(self):
        h = {"Authorization": f"Bearer {self.token}"}
        # sin encargos pendientes para esa fecha no hay nada que imprimir
        r = self.c.post("/api/encargos/hoja-entrega", headers=h, json={"fecha": "2031-01-01"})
        self.assertEqual(r.status_code, 400)
        a1, a2 = self._pieza("Hoja A", 5, 30, 90), self._pieza("Hoja B", 2, 15, 45)
        for cliente, cant in (("Ana Balderas", 2), ("Beto Balderas", 3)):
            r = self.c.post("/api/encargos", headers=h, json={
                "cliente_nuevo": cliente, "fecha_entrega": "2031-02-01", "notas": "Empacar en bolsa",
                "anticipo": 30, "forma_anticipo": "Depósito",
                "items": [{"articulo_id": a1["id"], "cantidad": cant}, {"articulo_id": a2["id"], "cantidad": 1}]})
            self.assertEqual(r.status_code, 201, r.get_json())
        for cfg in ({"fecha": "2031-02-01"}, {}, {"fecha": "2031-02-01", "incluir_costos": False}):
            r = self.c.post("/api/encargos/hoja-entrega", headers=h, json=cfg)
            self.assertEqual(r.status_code, 200, cfg)
            self.assertTrue(r.data.startswith(b"%PDF"))
        self.assertEqual(self.c.post("/api/encargos/hoja-entrega", json={}).status_code, 401)
        self.assertEqual(self.c.post("/api/encargos/hoja-entrega", headers={"Authorization": f"Bearer {self.token2}"},
                                     json={"fecha": "2031-02-01"}).status_code, 400, "otra cuenta no ve estos encargos")

    def test_38_apartado_lleva_lugar_de_entrega(self):
        h = {"Authorization": f"Bearer {self.token}"}
        a = self._pieza("Apartado lugar", 3, 10, 50)
        r = self.c.post("/api/apartados", headers=h, json={"articulo_id": a["id"], "precio_acordado": 50, "anticipo": 10, "cliente_nuevo": "Cliente 38"})
        self.assertEqual(r.status_code, 201)
        self.assertEqual(r.get_json()["lugar_entrega"], "Balderas", "por defecto se entrega en Balderas")
        r = self.c.post("/api/apartados", headers=h, json={"articulo_id": a["id"], "precio_acordado": 50, "lugar_entrega": "Punto de encuentro", "cliente_nuevo": "Cliente 38b"})
        self.assertEqual(r.get_json()["lugar_entrega"], "Punto de encuentro")

    def test_39_migracion_agrega_lugar_de_entrega_a_bases_viejas(self):
        import sqlite3
        from collecthub import db as bd_mod
        viejo = Path(tempfile.gettempdir()) / f"ch-vieja-{os.getpid()}.db"
        sql = (Path(bd_mod.__file__).parent / "schema.sql").read_text(encoding="utf-8").replace(
            "  lugar_entrega   TEXT NOT NULL DEFAULT 'Balderas',\n", "")
        con = sqlite3.connect(viejo)
        con.executescript(sql)
        con.close()
        try:
            with patch.object(bd_mod, "RUTA_BD", viejo):
                bd_mod.crear_esquema()
                bd_mod.crear_esquema()   # idempotente
            con = sqlite3.connect(viejo)
            try:
                self.assertIn("lugar_entrega", [r[1] for r in con.execute("PRAGMA table_info(apartados)")])
            finally:
                con.close()
        finally:
            for suf in ("", "-wal", "-shm"):
                Path(str(viejo) + suf).unlink(missing_ok=True)

    def test_40_todo_apartado_lleva_el_nombre_del_cliente(self):
        h = {"Authorization": f"Bearer {self.token}"}
        a = self._pieza("Apartado cliente", 5, 10, 50)
        base = {"articulo_id": a["id"], "precio_acordado": 50, "anticipo": 10}
        self.assertEqual(self.c.post("/api/apartados", headers=h, json=base).status_code, 400, "sin cliente no se aparta")
        self.assertEqual(self.c.post("/api/apartados", headers=h, json={**base, "comprador_id": "C-NOEXISTE"}).status_code, 404)
        r = self.c.post("/api/apartados", headers=h, json={**base, "cliente_nuevo": "Marisol Tianguis", "tel_nuevo": "555"})
        self.assertEqual(r.status_code, 201, r.get_json())
        ap = r.get_json()
        self.assertEqual(ap["cliente_snap"], "Marisol Tianguis")
        _, est = self.pedir("GET", "/api/estado")
        cliente = next(c for c in est["compradores"] if c["id"] == ap["comprador_id"])
        self.assertEqual((cliente["nombre"], cliente["tel"]), ("Marisol Tianguis", "555"), "el cliente nuevo se da de alta")
        r2 = self.c.post("/api/apartados", headers=h, json={**base, "comprador_id": cliente["id"]})
        self.assertEqual(r2.get_json()["cliente_snap"], "Marisol Tianguis")
        # si después borran al cliente, el apartado conserva su nombre
        self.assertEqual(self.c.delete(f"/api/compradores/{cliente['id']}", headers=h).status_code, 200)
        _, est = self.pedir("GET", "/api/estado")
        viejo = next(x for x in est["apartados"] if x["id"] == ap["id"])
        self.assertEqual((viejo["comprador_id"], viejo["cliente_snap"]), (None, "Marisol Tianguis"))
        # un cliente de otra cuenta no se puede usar
        _, est2 = self.pedir("GET", "/api/estado", token=self.token2)
        if est2["compradores"]:
            self.assertEqual(self.c.post("/api/apartados", headers=h, json={**base, "comprador_id": est2["compradores"][0]["id"]}).status_code, 404)

    def test_41_ids_no_se_repiten_ni_en_el_mismo_milisegundo(self):
        from collecthub.util import uid
        with patch("collecthub.util.time.time", return_value=1_800_000_000.123):   # el tiempo "se congela"
            ids = [uid("HW") for _ in range(3000)]
        self.assertEqual(len(set(ids)), 3000)
        self.assertRegex(ids[0], r"^HW-[0-9A-Z]{11}$")

    def test_42_hoja_de_entrega_en_excel(self):
        import openpyxl
        h = {"Authorization": f"Bearer {self.token}"}
        self.assertEqual(self.c.post("/api/encargos/hoja-entrega-excel", headers=h, json={"fecha": "2032-01-01"}).status_code, 400)
        a1, a2 = self._pieza("Excel A", 5, 30, 90), self._pieza("Excel B", 2, 15, 45)
        for cliente, cant in (("Ana Excel", 2), ("Beto Excel", 3)):
            r = self.c.post("/api/encargos", headers=h, json={
                "cliente_nuevo": cliente, "fecha_entrega": "2032-02-07", "anticipo": 30, "forma_anticipo": "Depósito",
                "items": [{"articulo_id": a1["id"], "cantidad": cant}, {"articulo_id": a2["id"], "cantidad": 1}]})
            self.assertEqual(r.status_code, 201, r.get_json())
        r = self.c.post("/api/encargos/hoja-entrega-excel", headers=h, json={"fecha": "2032-02-07"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.mimetype)
        wb = openpyxl.load_workbook(io.BytesIO(r.data))
        self.assertEqual(wb.sheetnames, ["Pedidos", "Por cliente", "Empacar"])
        ws = wb["Pedidos"]
        filas = [f for f in ws.iter_rows(min_row=2, values_only=True) if f[1]]
        self.assertEqual(len(filas), 4, "una fila por pieza: 2 clientes x 2 piezas")
        cab = [c.value for c in ws[1]]
        self.assertIn("Costo total", cab)
        # el anticipo del pedido va solo en su primera fila (así no se duplica al sumar)
        col_ant = cab.index("Anticipo (del pedido)")
        self.assertEqual(sum(1 for f in filas if f[col_ant]), 2)
        self.assertEqual(wb["Empacar"]["D2"].value + wb["Empacar"]["D3"].value, 7)
        # sin costos no aparecen ni costo ni ganancia
        r = self.c.post("/api/encargos/hoja-entrega-excel", headers=h, json={"fecha": "2032-02-07", "incluir_costos": False})
        cab = [c.value for c in openpyxl.load_workbook(io.BytesIO(r.data))["Pedidos"][1]]
        self.assertNotIn("Costo total", cab)
        self.assertNotIn("Ganancia", cab)
        self.assertEqual(self.c.post("/api/encargos/hoja-entrega-excel", json={}).status_code, 401)
        self.assertEqual(self.c.post("/api/encargos/hoja-entrega-excel", headers={"Authorization": f"Bearer {self.token2}"},
                                     json={"fecha": "2032-02-07"}).status_code, 400, "otra cuenta no ve estos pedidos")

    def test_43_cobranza_abonos_despues_de_entregar(self):
        h = {"Authorization": f"Bearer {self.token}"}
        a = self._pieza("Cobranza", 2, 40, 100)
        e = self.c.post("/api/encargos", headers=h, json={"cliente_nuevo": "Deudor", "anticipo": 20, "forma_anticipo": "Efectivo",
                        "items": [{"articulo_id": a["id"], "cantidad": 2}]}).get_json()
        self.assertEqual(self.c.post(f"/api/encargos/{e['id']}/pago", headers=h, json={"monto": 10}).status_code, 409, "aún no entregado")
        d = self.c.post(f"/api/encargos/{e['id']}/entregar", headers=h, json={"cobrado": 100, "forma": "Efectivo"}).get_json()
        self.assertEqual((d["debe"], d["estatus"]), (80.0, "Entregado"), "200 - 20 anticipo - 100 cobrado")
        self.assertEqual(self.c.post(f"/api/encargos/{e['id']}/pago", headers=h, json={"monto": 0}).status_code, 400)
        self.assertEqual(self.c.post(f"/api/encargos/{e['id']}/pago", headers=h, json={"monto": 999}).status_code, 409, "no puede abonar de más")
        self.assertEqual(self.c.post(f"/api/encargos/{e['id']}/pago", headers={"Authorization": f"Bearer {self.token2}"}, json={"monto": 5}).status_code, 404)
        r = self.c.post(f"/api/encargos/{e['id']}/pago", headers=h, json={"monto": 30, "forma": "Depósito"})
        self.assertEqual((r.status_code, r.get_json()["debe"]), (201, 50.0))
        r = self.c.post(f"/api/encargos/{e['id']}/pago", headers=h, json={"monto": 50, "forma": "Efectivo"})
        self.assertEqual(r.get_json()["debe"], 0.0)
        self.assertEqual(len(r.get_json()["pagos"]), 2)
        self.assertEqual(self.c.post(f"/api/encargos/{e['id']}/pago", headers=h, json={"monto": 1}).status_code, 409, "ya liquidado")

    def test_44_catalogo_pdf_con_fotos(self):
        from unittest.mock import MagicMock
        h = {"Authorization": f"Bearer {self.token}"}
        self.assertEqual(self.c.post("/api/articulos/catalogo-pdf", json={}).status_code, 401)
        a = self._pieza("Catalogo con foto", 1, 10, 55)
        self.c.post(f"/api/articulos/{a['id']}/foto", data={"foto": (self._imagen(), "f.jpg")}, headers=h, content_type="multipart/form-data")
        _, ext = self.pedir("POST", "/api/articulos", {"tipo": "Pokémon", "nombre": "Carta con URL ajena", "cantidad": 1, "valor_estimado": 20,
                                                     "foto": "https://evil.example.com/foto.jpg"})
        for i in range(9):
            self._pieza(f"Catalogo relleno {i}", 1, 5, 25)
        with patch("collecthub.catalogo.urllib.request.urlopen", side_effect=AssertionError("no debe descargar de hosts ajenos")):
            r = self.c.post("/api/articulos/catalogo-pdf", headers=h, json={})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.data.startswith(b"%PDF"))
        paginas = r.data.count(b"/Type /Page") - r.data.count(b"/Type /Pages")
        self.assertGreaterEqual(paginas, 2, "más de 9 piezas ocupan más de una página")
        # una selección concreta
        r = self.c.post("/api/articulos/catalogo-pdf", headers=h, json={"ids": [a["id"]]})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self.c.post("/api/articulos/catalogo-pdf", headers=h, json={"ids": ["HW-NOEXISTE"]}).status_code, 400)
        # el catálogo no incluye piezas 'Conservar' ni 'Por recibir'
        _, cons = self.pedir("POST", "/api/articulos", {"tipo": "Hot Wheels", "nombre": "Solo para mi", "estatus": "Conservar", "valor_estimado": 99})
        self.assertEqual(self.c.post("/api/articulos/catalogo-pdf", headers=h, json={"ids": [cons["id"]]}).status_code, 400)


if __name__ == "__main__":
    unittest.main(verbosity=2)
