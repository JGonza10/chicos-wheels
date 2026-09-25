"""Hoja de entrega en Excel (.xlsx): la misma información que el PDF, pero como lista
de hoja de cálculo (una fila por pieza, filtros, totales con fórmulas) para ordenar,
filtrar y hacer cuentas.

Hojas: "Pedidos" (una fila por pieza), "Por cliente" (una fila por pedido, con columnas
para anotar lo cobrado y la diferencia) y "Empacar" (piezas juntas por ubicación).
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .db import todos
from .reporte import _fecha_ok
from .util import ErrorApp

ETIQUETA = {"Pendiente": "Apartado", "Empacado": "En proceso", "Entregado": "Liquidado", "Cancelado": "Cancelado"}
MONEDA = '"$"#,##0.00'
AZUL = "0A1F44"
_borde = Side(style="thin", color="C8CED8")
BORDE = Border(left=_borde, right=_borde, top=_borde, bottom=_borde)


def _encabezado(ws, columnas):
    ws.append([c[0] for c in columnas])
    for i, (_, ancho, _fmt) in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def _formatos(ws, columnas, desde, hasta):
    for fila in ws.iter_rows(min_row=desde, max_row=hasta):
        for celda, (_, _, fmt) in zip(fila, columnas):
            celda.border = BORDE
            if fmt:
                celda.number_format = fmt


def _fila_total(ws, fila, columnas, sumar, desde, hasta, etiqueta_col=1):
    ws.cell(row=fila, column=etiqueta_col, value="TOTAL")
    for idx in sumar:
        letra = get_column_letter(idx)
        ws.cell(row=fila, column=idx, value=f"=SUM({letra}{desde}:{letra}{hasta})")
    for i, (_, _, fmt) in enumerate(columnas, start=1):
        c = ws.cell(row=fila, column=i)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E1E6F0")
        c.border = BORDE
        if fmt and i in sumar:
            c.number_format = fmt


def generar_hoja_excel(usuario_id: str, config: dict) -> bytes:
    from .rutas.movimientos import ACTIVOS, encargos_de

    fecha_f = _fecha_ok(config.get("fecha"))
    incluir_costos = config.get("incluir_costos", True) is not False
    todos_estatus = config.get("incluir_entregados") is True
    encargos = [e for e in encargos_de(usuario_id)
                if (e["estatus"] in ACTIVOS or (todos_estatus and e["estatus"] == "Entregado"))
                and (not fecha_f or e["fecha_entrega"] == fecha_f)]
    if not encargos:
        raise ErrorApp("No hay pedidos" + (f" para el {fecha_f}" if fecha_f else " pendientes") + " que exportar")
    clientes = {c["id"]: c for c in todos("SELECT id,nombre,tel FROM compradores WHERE usuario_id=?", (usuario_id,))}
    arts = {a["id"]: a for a in todos("SELECT id,ubicacion,numero FROM articulos WHERE usuario_id=?", (usuario_id,))}
    encargos.sort(key=lambda e: (e["fecha_entrega"] or "9999", (clientes.get(e["comprador_id"]) or {}).get("nombre", "").lower()))

    wb = Workbook()

    # ---------------- Hoja 1: una fila por pieza ----------------
    ws = wb.active
    ws.title = "Pedidos"
    cols = [("Fecha entrega", 14, None), ("Cliente", 22, None), ("Teléfono", 15, None), ("Estatus", 12, None),
            ("Pieza", 38, None), ("Núm.", 10, None), ("Cant.", 7, "0")]
    if incluir_costos:
        cols += [("Costo c/u", 12, MONEDA), ("Costo total", 13, MONEDA)]
    cols += [("Precio c/u", 12, MONEDA), ("Precio total", 13, MONEDA)]
    if incluir_costos:
        cols += [("Ganancia", 13, MONEDA)]
    cols += [("Anticipo (del pedido)", 14, MONEDA), ("Forma anticipo", 13, None), ("Resta por cobrar (del pedido)", 16, MONEDA),
             ("Entregado ✔", 11, None), ("Notas", 30, None)]
    _encabezado(ws, cols)
    pos = {c[0]: i for i, c in enumerate(cols, start=1)}
    fila = 2
    for e in encargos:
        c = clientes.get(e["comprador_id"]) or {}
        for n, it in enumerate(e["items"]):
            base = arts.get(it["articulo_id"]) or {}
            f = fila
            valores = {
                "Fecha entrega": e["fecha_entrega"] or "", "Cliente": c.get("nombre", "Cliente"), "Teléfono": c.get("tel", ""),
                "Estatus": ETIQUETA.get(e["estatus"], e["estatus"]), "Pieza": it["nombre_snap"], "Núm.": base.get("numero", ""),
                "Cant.": it["cantidad"], "Precio c/u": it["precio_unit"],
                # el anticipo y lo que resta son del pedido: solo en su primera fila para que las sumas no se dupliquen
                "Anticipo (del pedido)": e["anticipo"] if n == 0 else None,
                "Forma anticipo": (e["forma_anticipo"] if n == 0 else ""),
                "Resta por cobrar (del pedido)": e["resta"] if n == 0 else None,
                "Entregado ✔": "", "Notas": e["notas"] if n == 0 else "",
            }
            if incluir_costos:
                valores["Costo c/u"] = it["costo_unit"]
            for k, v in valores.items():
                ws.cell(row=f, column=pos[k], value=v)
            letra_cant = get_column_letter(pos["Cant."])
            letra_pre = get_column_letter(pos["Precio c/u"])
            ws.cell(row=f, column=pos["Precio total"], value=f"={letra_cant}{f}*{letra_pre}{f}")
            if incluir_costos:
                letra_cos = get_column_letter(pos["Costo c/u"])
                ws.cell(row=f, column=pos["Costo total"], value=f"={letra_cant}{f}*{letra_cos}{f}")
                ws.cell(row=f, column=pos["Ganancia"],
                        value=f"={get_column_letter(pos['Precio total'])}{f}-{get_column_letter(pos['Costo total'])}{f}")
            fila += 1
    ultimo = fila - 1
    _formatos(ws, cols, 2, ultimo)
    sumar = [pos[k] for k in ("Cant.", "Costo total", "Precio total", "Ganancia", "Anticipo (del pedido)", "Resta por cobrar (del pedido)") if k in pos]
    _fila_total(ws, fila, cols, sumar, 2, ultimo, etiqueta_col=pos["Pieza"])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ultimo}"

    # ---------------- Hoja 2: una fila por pedido (cruce contra dinero) ----------------
    w2 = wb.create_sheet("Por cliente")
    cols2 = [("Fecha entrega", 14, None), ("Cliente", 24, None), ("Teléfono", 15, None), ("Estatus", 12, None), ("Piezas", 8, "0")]
    if incluir_costos:
        cols2 += [("Costo", 13, MONEDA)]
    cols2 += [("Precio", 13, MONEDA)]
    if incluir_costos:
        cols2 += [("Ganancia", 13, MONEDA)]
    cols2 += [("Anticipo", 13, MONEDA), ("Forma anticipo", 13, None), ("Por cobrar", 14, MONEDA),
              ("Cobrado efectivo", 15, MONEDA), ("Cobrado depósito", 15, MONEDA), ("Diferencia", 13, MONEDA), ("Firma / notas", 26, None)]
    _encabezado(w2, cols2)
    p2 = {c[0]: i for i, c in enumerate(cols2, start=1)}
    for i, e in enumerate(encargos, start=2):
        c = clientes.get(e["comprador_id"]) or {}
        vals = {"Fecha entrega": e["fecha_entrega"] or "", "Cliente": c.get("nombre", "Cliente"), "Teléfono": c.get("tel", ""),
                "Estatus": ETIQUETA.get(e["estatus"], e["estatus"]), "Piezas": e["piezas"], "Precio": e["total"],
                "Anticipo": e["anticipo"], "Forma anticipo": e["forma_anticipo"], "Por cobrar": e["resta"]}
        if incluir_costos:
            vals["Costo"] = e["costo"]
            vals["Ganancia"] = e["ganancia"]
        for k, v in vals.items():
            w2.cell(row=i, column=p2[k], value=v)
        ef, dep, por = (get_column_letter(p2[k]) for k in ("Cobrado efectivo", "Cobrado depósito", "Por cobrar"))
        # Diferencia = lo que cobraste menos lo que debías cobrar (debe quedar en $0)
        w2.cell(row=i, column=p2["Diferencia"], value=f"=({ef}{i}+{dep}{i})-{por}{i}")
    u2 = len(encargos) + 1
    _formatos(w2, cols2, 2, u2)
    sumar2 = [p2[k] for k in ("Piezas", "Costo", "Precio", "Ganancia", "Anticipo", "Por cobrar", "Cobrado efectivo", "Cobrado depósito", "Diferencia") if k in p2]
    _fila_total(w2, u2 + 1, cols2, sumar2, 2, u2, etiqueta_col=p2["Cliente"])
    w2.auto_filter.ref = f"A1:{get_column_letter(len(cols2))}{u2}"

    # ---------------- Hoja 3: lista para empacar ----------------
    w3 = wb.create_sheet("Empacar")
    cols3 = [("Ubicación", 24, None), ("Pieza", 40, None), ("Núm.", 10, None), ("Cant.", 8, "0"), ("Empacada ✔", 12, None)]
    _encabezado(w3, cols3)
    consol = {}
    for e in encargos:
        for it in e["items"]:
            base = arts.get(it["articulo_id"]) or {}
            k = it["articulo_id"] or it["nombre_snap"]
            x = consol.setdefault(k, {"ubic": base.get("ubicacion") or "-", "nombre": it["nombre_snap"], "num": base.get("numero", ""), "cant": 0})
            x["cant"] += it["cantidad"]
    for i, x in enumerate(sorted(consol.values(), key=lambda x: (x["ubic"], x["nombre"])), start=2):
        w3.cell(row=i, column=1, value=x["ubic"])
        w3.cell(row=i, column=2, value=x["nombre"])
        w3.cell(row=i, column=3, value=x["num"])
        w3.cell(row=i, column=4, value=x["cant"])
    u3 = len(consol) + 1
    _formatos(w3, cols3, 2, u3)
    _fila_total(w3, u3 + 1, cols3, [4], 2, u3, etiqueta_col=2)
    w3.auto_filter.ref = f"A1:E{u3}"

    for hoja in wb.worksheets:
        hoja.sheet_view.zoomScale = 110
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.fitToWidth = 1
        hoja.page_setup.fitToHeight = 0
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
        hoja.print_title_rows = "1:1"

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
